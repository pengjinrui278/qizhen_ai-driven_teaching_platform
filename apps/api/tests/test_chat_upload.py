"""对话里的文档上传：抽正文、按账号隔离、随账号注销清理。

边界也在这里钉住：不支持的格式拒收、抽不出文字要如实说、不能假装读到了。
"""
import io

import pytest
from fastapi.testclient import TestClient

from mirror_api import chat_upload
from mirror_api.chat_upload import ChatUpload, extract_text, load_uploads
from mirror_api.main import app, configure, _auth_attempts

H = {"X-Mirror-Request": "1"}


@pytest.fixture()
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("MIRROR_STAFF_INVITE_CODE", "staff-test")
    monkeypatch.setattr(chat_upload, "UPLOAD_ROOT", tmp_path / "chat-uploads")
    configure(app, "sqlite:///" + str(tmp_path / "uploads.sqlite"))
    _auth_attempts.clear()
    result = TestClient(app, headers=H)
    register(result, "upload_student")
    yield result
    app.state.engine.dispose()


def register(client, name, role="student"):
    r = client.post("/api/v2/auth/register", json={"username": name, "password": "test-only-password",
                                                   "nickname": name, "role": role, "invite_code": "staff-test"})
    assert r.status_code == 200, r.text
    return r.json()


def docx_bytes(*paragraphs):
    import docx
    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def pptx_bytes(*lines):
    from pptx import Presentation
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    box = slide.shapes.add_textbox(0, 0, 100, 100)
    box.text_frame.text = "\n".join(lines)
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def test_authentication_required(client):
    anonymous = TestClient(app, headers=H)
    assert anonymous.post("/api/v2/uploads",
                          files={"file": ("a.txt", b"hi", "text/plain")}).status_code == 401


def test_text_upload_is_readable(client):
    made = client.post("/api/v2/uploads",
                       files={"file": ("笔记.txt", "反常积分的比较判别法".encode("utf-8"), "text/plain")})
    assert made.status_code == 200, made.text
    body = made.json()
    assert body["kind"] == "text" and body["readable"] is True
    assert body["chars"] > 0 and "note" not in body

    with app.state.session_factory() as db:
        rows = load_uploads(db, client.get("/api/v2/me").json()["id"], [body["id"]])
        assert "比较判别法" in rows[0].text, "正文要存下来，agent 才读得到"


def test_word_and_ppt_are_extracted(client):
    word = client.post("/api/v2/uploads",
                       files={"file": ("讲义.docx", docx_bytes("定义 8.1 无穷限积分", "定理 8.1 比较判别法"),
                                       "application/vnd.openxmlformats-officedocument.wordprocessingml.document")})
    assert word.status_code == 200, word.text
    assert word.json()["kind"] == "word" and word.json()["readable"] is True

    deck = client.post("/api/v2/uploads",
                       files={"file": ("课件.pptx", pptx_bytes("第8章 反常积分", "判别法小结"),
                                       "application/vnd.openxmlformats-officedocument.presentationml.presentation")})
    assert deck.status_code == 200, deck.text
    assert deck.json()["kind"] == "ppt" and deck.json()["readable"] is True

    with app.state.session_factory() as db:
        account = client.get("/api/v2/me").json()["id"]
        texts = {row.filename: row.text for row in load_uploads(db, account, [word.json()["id"], deck.json()["id"]])}
        assert "比较判别法" in texts["讲义.docx"]
        assert "第8章 反常积分" in texts["课件.pptx"]


def test_scanned_pdf_says_it_could_not_be_read(client):
    """抽不出文字时必须如实说，不能让模型以为读到了。"""
    payload = b"%PDF-1.4\n1 0 obj<</Type/Catalog>>endobj\ntrailer<</Root 1 0 R>>\n%%EOF"
    made = client.post("/api/v2/uploads", files={"file": ("扫描件.pdf", payload, "application/pdf")})
    assert made.status_code == 200, made.text
    body = made.json()
    assert body["readable"] is False
    assert "读不到" in body.get("note", ""), "要明确说读不到，而不是静默返回空"


def test_rejects_unsupported_and_empty(client):
    assert client.post("/api/v2/uploads",
                       files={"file": ("x.exe", b"MZ", "application/octet-stream")}).status_code == 415
    assert client.post("/api/v2/uploads",
                       files={"file": ("x.txt", b"", "text/plain")}).status_code == 400


def test_uploads_are_private_to_the_account(client):
    made = client.post("/api/v2/uploads",
                       files={"file": ("a.txt", b"only mine", "text/plain")}).json()
    other = TestClient(app, headers=H)
    register(other, "upload_teacher", "teacher")
    with app.state.session_factory() as db:
        assert load_uploads(db, other.get("/api/v2/me").json()["id"], [made["id"]]) == []

    exported = other.get("/api/v2/me/export").json()
    assert exported["chat_uploads"] == [], "别人的导出里不该出现我的上传"


def test_delete_account_purges_uploads_with_files(client):
    made = client.post("/api/v2/uploads",
                       files={"file": ("a.txt", b"data", "text/plain")}).json()
    with app.state.session_factory() as db:
        stored = db.get(ChatUpload, made["id"]).storage_key
        path = chat_upload.UPLOAD_ROOT / stored
        assert path.is_file()
        exported = client.get("/api/v2/me/export").json()
        assert exported["chat_uploads"][0]["filename"] == "a.txt"

    assert client.post("/api/v2/me/delete", json={"password": "test-only-password",
                                                  "confirmed": True}).status_code == 200
    with app.state.session_factory() as db:
        assert db.query(ChatUpload).count() == 0
    assert not path.exists(), "注销时磁盘原件也要删掉"


def test_extract_text_handles_broken_files_without_raising():
    """坏文件按「读不出来」处理，不能让上传整个失败。"""
    assert extract_text("坏.pdf", b"not a pdf at all", "pdf") == ""
    assert extract_text("坏.docx", b"not a docx", "word") == ""
    # 老格式 .doc/.ppt 没有纯 Python 解析器 → 如实返回空
    assert extract_text("老.doc", b"\xd0\xcf\x11\xe0", "word") == ""
